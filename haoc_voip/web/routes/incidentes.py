"""
Rotas do Módulo de Incidentes e Histórico Operacional.
Filtros por ramal, bloco, período, criticidade e status; edição de causa/observação; exportação CSV e Excel.
"""
from __future__ import annotations

import csv
import io
import logging
from datetime import datetime
from flask import Blueprint, render_template, request, redirect, url_for, flash, Response, jsonify, session

from haoc_voip.web.app import login_required, analista_ou_admin_required
from haoc_voip.core.models import Incidente, Ramal, Bloco
from haoc_voip.core.database import db
from haoc_voip.core.audit import registrar_auditoria

logger = logging.getLogger("haoc.routes.incidentes")
incidentes_bp = Blueprint("incidentes", __name__)


@incidentes_bp.route("/")
@login_required
def lista():
    """Exibe o histórico de incidentes com filtros dinâmicos."""
    ramal_id = request.args.get("ramal_id", type=int)
    bloco = request.args.get("bloco", "").strip()
    status_filtro = request.args.get("status", "").strip()  # aberto, encerrado
    criticidade = request.args.get("criticidade", "").strip()
    data_inicio_str = request.args.get("data_inicio", "").strip()
    data_fim_str = request.args.get("data_fim", "").strip()

    with db.session_scope() as session_db:
        query = session_db.query(Incidente).join(Ramal, Incidente.ramal_id == Ramal.id)

        if ramal_id:
            query = query.filter(Incidente.ramal_id == ramal_id)
        if bloco:
            query = query.filter(Ramal.bloco == bloco)
        if criticidade:
            query = query.filter(Ramal.criticidade == criticidade)
        if status_filtro == "aberto":
            query = query.filter(Incidente.encerrado_em == None)
        elif status_filtro == "encerrado":
            query = query.filter(Incidente.encerrado_em != None)

        if data_inicio_str:
            try:
                dt_ini = datetime.strptime(data_inicio_str, "%Y-%m-%d")
                query = query.filter(Incidente.iniciado_em >= dt_ini)
            except ValueError:
                pass

        if data_fim_str:
            try:
                dt_fim = datetime.strptime(data_fim_str + " 23:59:59", "%Y-%m-%d %H:%M:%S")
                query = query.filter(Incidente.iniciado_em <= dt_fim)
            except ValueError:
                pass

        incidentes = query.order_by(Incidente.iniciado_em.desc()).limit(150).all()
        incidentes_dados = [i.to_dict() for i in incidentes]

        blocos = session_db.query(Bloco).filter(Bloco.ativo == True).order_by(Bloco.nome).all()
        ramais = session_db.query(Ramal).filter(Ramal.ativo == True).order_by(Ramal.descricao).all()

        return render_template(
            "incidentes.html",
            incidentes=incidentes_dados,
            blocos=[b.to_dict() for b in blocos],
            ramais=[r.to_dict() for r in ramais],
            filtros={
                "ramal_id": ramal_id,
                "bloco": bloco,
                "status": status_filtro,
                "criticidade": criticidade,
                "data_inicio": data_inicio_str,
                "data_fim": data_fim_str,
            },
        )


@incidentes_bp.route("/<int:incidente_id>/observacao", methods=["POST"])
@analista_ou_admin_required
def atualizar_observacao(incidente_id: int):
    """Permite ao analista registrar ou alterar a causa investigada e observações."""
    causa = request.form.get("causa", "").strip()
    observacao = request.form.get("observacao", "").strip()

    with db.session_scope() as session_db:
        inc = session_db.query(Incidente).filter(Incidente.id == incidente_id).first()
        if not inc:
            flash("Incidente não encontrado.", "danger")
            return redirect(url_for("incidentes.lista"))

        inc.causa = causa or "Indisponibilidade de rede"
        inc.observacao = observacao or None

    registrar_auditoria(
        acao="ATUALIZAR_INCIDENTE",
        entidade="incidente",
        entidade_id=incidente_id,
        detalhes={"causa": causa, "observacao": observacao},
        usuario_id=session.get("usuario_id"),
        ip_origem=request.remote_addr,
    )

    flash(f"Observações do incidente #{incidente_id} atualizadas.", "success")
    return redirect(request.referrer or url_for("incidentes.lista"))


@incidentes_bp.route("/exportar/csv")
@login_required
def exportar_csv():
    """Gera arquivo CSV para download com encoding UTF-8 com BOM para Excel."""
    output = io.StringIO()
    # Adiciona UTF-8 BOM para acentuação correta no Microsoft Excel
    output.write("\ufeff")
    writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_MINIMAL)

    writer.writerow([
        "ID Incidente",
        "Ramal",
        "Bloco",
        "IP Monitorado",
        "MAC Cisco",
        "Início da Queda",
        "Retorno",
        "Duração (segundos)",
        "Duração Formatada",
        "Status",
        "Causa Investigada",
        "Observação do Analista",
    ])

    with db.session_scope() as session_db:
        incidentes = session_db.query(Incidente).order_by(Incidente.iniciado_em.desc()).all()
        for inc in incidentes:
            d = inc.to_dict()
            writer.writerow([
                d["id"],
                d["ramal_descricao"],
                d["bloco"],
                d["ip_monitorado"],
                d["mac_cisco"],
                d["iniciado_em"],
                d["encerrado_em"],
                d["duracao_segundos"] or 0,
                d["duracao_formatada"],
                "ABERTO" if d["aberto"] else "RESOLVIDO",
                d["causa"],
                d["observacao"],
            ])

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Response(
        output.getvalue(),
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename=incidentes_haoc_{timestamp}.csv"},
    )


@incidentes_bp.route("/exportar/excel")
@login_required
def exportar_excel():
    """Gera pasta de trabalho Excel (.xlsx) nativa ou CSV formatado."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Incidentes HAOC VoIP"

        # Cabeçalho estilizado
        headers = [
            "ID", "Ramal", "Bloco", "IP Monitorado", "MAC Cisco",
            "Início da Queda", "Retorno", "Duração", "Status", "Causa", "Observações"
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        for col_num, header_title in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        with db.session_scope() as session_db:
            incidentes = session_db.query(Incidente).order_by(Incidente.iniciado_em.desc()).all()
            for row_idx, inc in enumerate(incidentes, 2):
                d = inc.to_dict()
                ws.append([
                    d["id"],
                    d["ramal_descricao"],
                    d["bloco"],
                    d["ip_monitorado"],
                    d["mac_cisco"],
                    d["iniciado_em"],
                    d["encerrado_em"],
                    d["duracao_formatada"],
                    "ABERTO" if d["aberto"] else "RESOLVIDO",
                    d["causa"],
                    d["observacao"],
                ])

        # Ajuste de largura de colunas
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        excel_stream = io.BytesIO()
        wb.save(excel_stream)
        excel_stream.seek(0)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return Response(
            excel_stream.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=incidentes_haoc_{timestamp}.xlsx"},
        )

    except ImportError:
        # Fallback elegante para CSV caso openpyxl não esteja carregado
        return exportar_csv()
